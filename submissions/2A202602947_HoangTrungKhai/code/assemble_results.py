"""Merge actual screening/training/inference/test predictions into submission workbook."""
from __future__ import annotations

import argparse
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import eval as ev


def clear(ws):
    if ws.max_row > 1: ws.delete_rows(2, ws.max_row-1)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--workbook",default="../submissions/2A202602947_HoangTrungKhai/results.xlsx")
    ap.add_argument("--training",default="runs/training_summary.csv")
    ap.add_argument("--inference",default="runs/inference_summary.csv")
    ap.add_argument("--pred-dir",default="predictions")
    ap.add_argument("--out-dir",default="runs")
    a=ap.parse_args(); book=load_workbook(a.workbook)

    tr=pd.read_csv(a.training)
    tw=book["Training"]; clear(tw)
    base=tr[tr.exp_id.eq("T00")]
    base_f1=float(base.iloc[0].macro_f1_val) if len(base) else np.nan
    axes={"T00":("baseline","công thức nền"),"T01_frozen":("initialization","freeze backbone"),
          "T02_color":("augmentation","color jitter"),"T03_label_smoothing":("loss","label smoothing 0.1"),
          "T04_focal":("loss","focal gamma 2"),"T05_balanced_sampler":("sampler","balanced sampler"),
          "T06_mixup":("mixing","Mixup alpha 0.4"),"T07_combo":("combined","color + label smoothing + EMA")}
    for _,r in tr.iterrows():
        ax,diff=axes.get(r.exp_id,("other","see config.json"))
        val_file=Path(a.pred_dir)/f"{r.exp_id}_seed{int(r.seed)}_val.csv"
        rare=None
        if val_file.exists():
            pred_val=ev.read_pred(val_file)
            met_val=ev.compute_metrics(pred_val.y_true,pred_val.y_pred,pred_val.probs)
            rare=float(np.mean([met_val["f1"][0],met_val["f1"][7]]))
        tw.append([r.exp_id,r.get("backbone","convnext_tiny"),ax,diff,int(r.seed),float(r.macro_f1_val),float(r.top1_val),
                   float(r.macro_f1_val)-base_f1 if np.isfinite(base_f1) else None,rare,
                   f"best epoch {int(r.best_epoch)}; one screening seed"])

    iw=book["Inference"]; clear(iw); lw=book["Latency"]; clear(lw)
    inf=pd.read_csv(a.inference); i00=float(inf.loc[inf.exp_id.eq("I00_1view"),"p50_ms"].iloc[0])
    gpu=str(inf.device.iloc[0]) if "device" in inf else "unknown"
    for _,r in inf.iterrows():
        p50=r.get("p50_ms",np.nan); p95=r.get("p95_ms",np.nan); p99=r.get("p99_ms",np.nan)
        iw.append([r.exp_id,r.method,"validation-selected checkpoint",1 if r.exp_id in ("I00_1view","I03_temperature") else (2 if r.exp_id=="I01_hflip_prob" else 3),
                   float(r.macro_f1_val),float(r.top1_val),float(r.ece_val),p50,p95,p99,(1000/p50 if pd.notna(p50) else None),
                   (float(p50)/i00 if pd.notna(p50) else None)])
        if pd.notna(p50): lw.append([r.method,gpu,"fp32",1,224,"no",p50,p95,p99,1000/p50,"see runtime"])

    fw=book["Final"]; clear(fw); pcw=book["PerClass"]; clear(pcw)
    pred=Path(a.pred_dir); groups={}
    for exp in ("F01","T00"):
        files=sorted(pred.glob(f"{exp}_seed*_test.csv"))
        if len(files)<3: raise SystemExit(f"Thiếu >=3 file test cho {exp}: chỉ có {len(files)}")
        items=[ev.read_pred(f) for f in files]
        by_seed=[]
        for p in items:
            m=ev.compute_metrics(p.y_true,p.y_pred,p.probs)
            by_seed.append((p,m))
            val_file=pred/f"{exp}_seed{p.seed}_val.csv"
            val_pred=ev.read_pred(val_file)
            val_metrics=ev.compute_metrics(val_pred.y_true,val_pred.y_pred,val_pred.probs)
            fw.append([exp,"validation-selected recipe + " + ("selected inference" if exp=="F01" else "I00_1view"),p.seed,
                       val_metrics["macro_f1"],m["macro_f1"],m["top1"],m["ece"],None,None,
                       f"validation CSV: {val_file.name}"])
        groups[exp]=by_seed
        ms=np.array([x[1]["macro_f1"] for x in by_seed]); acc=np.array([x[1]["top1"] for x in by_seed])
        fw.append([exp,"mean ± sample std across seeds",None,None,None,None,None,
                   f"{ms.mean():.4f} ± {ms.std(ddof=1):.4f}",f"{acc.mean():.4f} ± {acc.std(ddof=1):.4f}","test set evaluated once per seed"])
    # Per-class test metrics averaged across seeds and confusion matrices for diagnosis.
    for exp,items in groups.items():
        for c,name in enumerate(ev.CLASS_NAMES):
            prec=np.mean([m["precision"][c] for _,m in items]); rec=np.mean([m["recall"][c] for _,m in items])
            f1=np.mean([m["f1"][c] for _,m in items]); support=int(np.mean([m["support"][c] for _,m in items]))
            pcw.append([exp,name,support,prec,rec,f1])
    # Store mean confusion matrices for final and baseline, using only saved predictions.
    import matplotlib.pyplot as plt
    out=Path(a.out_dir); out.mkdir(parents=True,exist_ok=True)
    for exp,items in groups.items():
        cm=np.mean([m["confusion"] for _,m in items],axis=0)
        fig,ax=plt.subplots(figsize=(8,7)); im=ax.imshow(cm,cmap="Blues"); fig.colorbar(im,ax=ax)
        ax.set_xticks(range(9),ev.CLASS_NAMES,rotation=60,ha="right"); ax.set_yticks(range(9),ev.CLASS_NAMES)
        ax.set_xlabel("Predicted"); ax.set_ylabel("True"); ax.set_title(f"{exp}: mean test confusion matrix")
        fig.tight_layout(); fig.savefig(out/f"{exp}_confusion_matrix.png",dpi=160); plt.close(fig)
        np.savetxt(out/f"{exp}_confusion_matrix.csv",cm,delimiter=",",fmt="%.3f")
    book.save(a.workbook)
    print(f"Updated {a.workbook}; test confusion matrices in {out}")


if __name__=="__main__": main()
