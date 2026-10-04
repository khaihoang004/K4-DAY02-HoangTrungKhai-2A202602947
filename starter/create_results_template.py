"""Create the required results workbook with empty, clearly labeled sheets."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SHEETS = {
    "Backbones": ["exp_id", "backbone", "pretrained_tag", "params_M", "GMAC", "resolution", "epochs", "seed", "macro_F1_val", "top1_val", "train_seconds_per_epoch", "latency_batch1_ms", "notes"],
    "Training": ["exp_id", "backbone", "changed_axis", "difference_from_T00", "seed", "macro_F1_val", "top1_val", "delta_vs_T00", "rare_class_F1", "notes"],
    "Inference": ["exp_id", "method", "model_checkpoint", "views_or_models_K", "macro_F1_val", "top1_val", "ECE_val", "p50_ms_batch1", "p95_ms_batch1", "p99_ms_batch1", "images_per_s", "relative_cost_vs_I00"],
    "Final": ["exp_id", "configuration", "seed", "macro_F1_val", "macro_F1_test", "top1_test", "ECE_test", "macro_F1_mean_std", "top1_mean_std", "notes"],
    "PerClass": ["configuration", "class", "test_support", "precision", "recall", "F1"],
    "Latency": ["configuration", "GPU", "dtype", "batch", "resolution", "BN_fused", "p50_ms", "p95_ms", "p99_ms", "images_per_s", "torch_version"],
    "Summary": ["rank_by_macro_F1_val", "exp_id", "backbone", "training_recipe", "inference_method", "macro_F1_val", "top1_val", "latency_p95_ms", "params_M", "notes"],
}


def create(path="submissions/2A202602947_HoangTrungKhai/results.xlsx"):
    book = Workbook(); book.remove(book.active)
    for name, columns in SHEETS.items():
        ws = book.create_sheet(name); ws.append(columns); ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        for col, title in enumerate(columns, 1): ws.column_dimensions[get_column_letter(col)].width = min(34, max(14, len(title)+2))
    target = Path(path); target.parent.mkdir(parents=True, exist_ok=True); book.save(target)
    return target


if __name__ == "__main__": create()
